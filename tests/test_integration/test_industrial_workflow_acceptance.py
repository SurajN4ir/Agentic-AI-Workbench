"""
tests/test_integration/test_industrial_workflow_acceptance.py
============================================================
Comprehensive Industrial Workflow Acceptance Test:
1. Verifies browser SPA delivery from FastAPI root.
2. Ingests genuine industrial test fixtures (.docx spec, .xlsx defect register, .png inspection gauge).
3. Executes private semantic vector retrieval with source citations.
4. Generates genuine deliverable report (.docx) and analysis workbook (.xlsx) into the session workspace.
5. Re-opens and parses generated deliverables using docx and openpyxl, verifying actual industrial parameters.
6. Downloads generated artifacts via authenticated endpoints.
7. Enforces strict multi-tenant boundary checks: other non-admin users receive 403 Forbidden.
8. Enforces path traversal rejection.
"""

import uuid
from pathlib import Path

import docx
import openpyxl
import pytest
from app.database import get_db
from app.identity.models import User
from app.identity.service import create_access_token, hash_password
from app.main import create_app
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import async_sessionmaker

pytestmark = pytest.mark.integration


@pytest.mark.asyncio
async def test_complete_industrial_workflow_acceptance(test_engine) -> None:
    app = create_app()
    session_factory = async_sessionmaker(test_engine, expire_on_commit=False)

    async def override_get_db():
        async with session_factory() as session:
            try:
                yield session
                await session.commit()
            except Exception:
                await session.rollback()
                raise

    app.dependency_overrides[get_db] = override_get_db

    # 1. Seed RBAC Users: Admin, Primary Analyst/Operator, and Rogue/Unauthorized Analyst
    admin_id = uuid.uuid4()
    admin_user = User(
        id=admin_id,
        username="lead_engineer",
        name="Lead Engineer",
        role="admin",
        password_hash=hash_password("LeadEngSecure2026!"),
        is_active=True,
    )

    operator_id = uuid.uuid4()
    operator_user = User(
        id=operator_id,
        username="operator_dan",
        name="Turbine Operator Dan",
        role="analyst",
        password_hash=hash_password("OperatorDanPass2026!"),
        is_active=True,
    )

    rogue_id = uuid.uuid4()
    rogue_user = User(
        id=rogue_id,
        username="unauthorized_tech",
        name="Unauthorized Tech",
        role="analyst",
        password_hash=hash_password("RogueTechPass2026!"),
        is_active=True,
    )

    async with session_factory() as session:
        session.add_all([admin_user, operator_user, rogue_user])
        await session.commit()

    admin_token = create_access_token(user_id=str(admin_id), role="admin", username="lead_engineer")
    operator_token = create_access_token(user_id=str(operator_id), role="analyst", username="operator_dan")
    rogue_token = create_access_token(user_id=str(rogue_id), role="analyst", username="unauthorized_tech")

    fixtures_dir = Path("workspace/industrial_fixtures")
    assert fixtures_dir.exists(), "Industrial fixtures directory must exist"
    spec_path = fixtures_dir / "turbine_maintenance_spec.docx"
    defects_path = fixtures_dir / "production_defects.xlsx"
    gauge_path = fixtures_dir / "test_gauge.png"

    assert spec_path.exists(), f"Missing {spec_path}"
    assert defects_path.exists(), f"Missing {defects_path}"
    assert gauge_path.exists(), f"Missing {gauge_path}"

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Step 1: Verify Browser SPA Delivery
        root_resp = await client.get("/")
        assert root_resp.status_code == 200
        assert "html" in root_resp.headers.get("content-type", "").lower()

        # Step 2: Upload genuine industrial test fixtures via authenticated knowledge API
        operator_headers = {"Authorization": f"Bearer {operator_token}"}

        # 2a. Ingest DOCX technical maintenance spec
        with open(spec_path, "rb") as f:
            docx_upload_resp = await client.post(
                "/api/v1/knowledge/upload",
                headers=operator_headers,
                files={"file": (spec_path.name, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"collection_id": "industrial_demo"},
            )
        assert docx_upload_resp.status_code == 200, docx_upload_resp.text
        docx_meta = docx_upload_resp.json()
        assert docx_meta["filename"] == "turbine_maintenance_spec.docx"
        assert docx_meta["file_type"] == ".docx"
        assert docx_meta["total_chunks"] >= 1

        # 2b. Ingest XLSX defect telemetry log
        with open(defects_path, "rb") as f:
            xlsx_upload_resp = await client.post(
                "/api/v1/knowledge/upload",
                headers=operator_headers,
                files={"file": (defects_path.name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                data={"collection_id": "industrial_demo"},
            )
        assert xlsx_upload_resp.status_code == 200, xlsx_upload_resp.text
        xlsx_meta = xlsx_upload_resp.json()
        assert xlsx_meta["filename"] == "production_defects.xlsx"
        assert xlsx_meta["file_type"] == ".xlsx"
        assert xlsx_meta["total_chunks"] >= 1

        # 2c. Ingest PNG inspection gauge image
        with open(gauge_path, "rb") as f:
            png_upload_resp = await client.post(
                "/api/v1/knowledge/upload",
                headers=operator_headers,
                files={"file": (gauge_path.name, f, "image/png")},
                data={"collection_id": "industrial_demo"},
            )
        assert png_upload_resp.status_code == 200, png_upload_resp.text
        png_meta = png_upload_resp.json()
        assert png_meta["filename"] == "test_gauge.png"
        assert png_meta["file_type"] == ".png"

        # Step 3: Semantic Vector Retrieval Verification
        search_resp = await client.post(
            "/api/v1/knowledge/search",
            headers=operator_headers,
            json={
                "query": "vibration velocity limit mm/s and emergency trip threshold",
                "collection_id": "industrial_demo",
                "top_k": 5,
            },
        )
        assert search_resp.status_code == 200, search_resp.text
        search_data = search_resp.json()
        assert search_data["results_count"] >= 1
        found_spec = any("4.5" in item["content"] or "turbine" in item["content"].lower() for item in search_data["results"])
        assert found_spec, "Search must retrieve turbine maintenance specification chunk"

        # Step 4: Deliverable Report Generation (.docx & .xlsx)
        session_id = str(uuid.uuid4())

        # 4a. Generate Word Report deliverable
        docx_req_payload = {
            "format": "docx",
            "session_id": session_id,
            "report_spec": {
                "title": "Industrial Turbine Health Assessment Report",
                "subtitle": "Confidential Plant Reliability Engineering Review",
                "author": "Operator Dan (Reliability Engineering)",
                "filename": "turbine_health_assessment.docx",
                "sections": [
                    {
                        "heading": "1. Operational Parameter Compliance",
                        "body": "According to turbine_maintenance_spec.docx, the radial vibration limit is 4.5 mm/s RMS. "
                                "Bearing thermal ceiling is 185 C and relief valve lifts at 150 PSI.",
                    },
                    {
                        "headers": ["Equipment ID", "Measured Vibration (mm/s)", "Thermal Ceiling (C)", "Status"],
                        "rows": [
                            ["Turbine-A101", "3.2", "162", "COMPLIANT"],
                            ["Turbine-B202", "5.1", "192", "CRITICAL_EXCEEDANCE"],
                            ["Turbine-C303", "4.1", "175", "MARGINAL"],
                        ],
                        "caption": "Table 1: Measured Fleet Telemetry vs. Engineering Thresholds",
                    },
                ],
            },
        }
        gen_docx_resp = await client.post(
            "/api/v1/artifacts/generate",
            headers=operator_headers,
            json=docx_req_payload,
        )
        assert gen_docx_resp.status_code == 201, gen_docx_resp.text
        docx_art_info = gen_docx_resp.json()
        assert docx_art_info["filename"] == "turbine_health_assessment.docx"
        assert docx_art_info["session_id"] == session_id

        # 4b. Generate Excel Workbook deliverable
        xlsx_req_payload = {
            "format": "xlsx",
            "session_id": session_id,
            "spreadsheet_spec": {
                "filename": "fleet_defect_metrics.xlsx",
                "sheets": [
                    {
                        "name": "DefectSummary",
                        "headers": ["Machine_ID", "Defect_Count", "Severity_Score", "Maintenance_Required"],
                        "rows": [
                            ["Turbine-A101", 1, 14.5, "NO"],
                            ["Turbine-B202", 4, 88.2, "IMMEDIATE"],
                            ["Turbine-C303", 2, 32.0, "SCHEDULED"],
                        ],
                    }
                ],
            },
        }
        gen_xlsx_resp = await client.post(
            "/api/v1/artifacts/generate",
            headers=operator_headers,
            json=xlsx_req_payload,
        )
        assert gen_xlsx_resp.status_code == 201, gen_xlsx_resp.text
        xlsx_art_info = gen_xlsx_resp.json()
        assert xlsx_art_info["filename"] == "fleet_defect_metrics.xlsx"

        # Step 5: Independent Verification of Generated Files on Disk
        gen_docx_path = Path("workspace") / session_id / "turbine_health_assessment.docx"
        gen_xlsx_path = Path("workspace") / session_id / "fleet_defect_metrics.xlsx"

        assert gen_docx_path.exists(), f"Word document not created on disk: {gen_docx_path}"
        assert gen_xlsx_path.exists(), f"Excel workbook not created on disk: {gen_xlsx_path}"

        # 5a. Inspect generated docx with python-docx
        parsed_doc = docx.Document(str(gen_docx_path))
        doc_paragraphs = [p.text for p in parsed_doc.paragraphs if p.text]
        assert any("Industrial Turbine Health Assessment" in t for t in doc_paragraphs)
        assert any("4.5 mm/s" in t for t in doc_paragraphs)
        assert len(parsed_doc.tables) == 1
        table = parsed_doc.tables[0]
        header_cells = [c.text.strip() for c in table.rows[0].cells]
        assert "Equipment ID" in header_cells
        assert "Status" in header_cells
        row_1_cells = [c.text.strip() for c in table.rows[1].cells]
        assert "Turbine-A101" in row_1_cells

        # 5b. Inspect generated xlsx with openpyxl
        parsed_wb = openpyxl.load_workbook(str(gen_xlsx_path))
        assert "DefectSummary" in parsed_wb.sheetnames
        ws = parsed_wb["DefectSummary"]
        headers = [cell.value for cell in ws[1]]
        assert headers == ["Machine_ID", "Defect_Count", "Severity_Score", "Maintenance_Required"]
        row2 = [cell.value for cell in ws[2]]
        assert row2[0] == "Turbine-A101"
        assert row2[3] == "NO"

        # Step 6: Authenticated Deliverable Download
        list_resp = await client.get(f"/api/v1/artifacts/{session_id}", headers=operator_headers)
        assert list_resp.status_code == 200
        artifacts_list = list_resp.json()
        assert len(artifacts_list) == 2
        file_names = [a["filename"] for a in artifacts_list]
        assert "turbine_health_assessment.docx" in file_names
        assert "fleet_defect_metrics.xlsx" in file_names

        # Download docx
        dl_docx_resp = await client.get(
            f"/api/v1/artifacts/{session_id}/turbine_health_assessment.docx",
            headers=operator_headers,
        )
        assert dl_docx_resp.status_code == 200
        assert "attachment" in dl_docx_resp.headers.get("content-disposition", "")
        assert len(dl_docx_resp.content) == gen_docx_path.stat().st_size

        # Download xlsx
        dl_xlsx_resp = await client.get(
            f"/api/v1/artifacts/{session_id}/fleet_defect_metrics.xlsx",
            headers=operator_headers,
        )
        assert dl_xlsx_resp.status_code == 200
        assert dl_xlsx_resp.content[:4] == b"PK\x03\x04"  # Valid ZIP / OpenXML header

        # Step 7: Multi-Tenant RBAC Boundary Enforcement
        rogue_headers = {"Authorization": f"Bearer {rogue_token}"}

        # Unauthorized user must receive 403 Forbidden attempting to list session artifacts
        rogue_list_resp = await client.get(f"/api/v1/artifacts/{session_id}", headers=rogue_headers)
        assert rogue_list_resp.status_code == 403
        assert "Access denied" in rogue_list_resp.text

        # Unauthorized user must receive 403 Forbidden attempting to download session artifact
        rogue_dl_resp = await client.get(
            f"/api/v1/artifacts/{session_id}/turbine_health_assessment.docx",
            headers=rogue_headers,
        )
        assert rogue_dl_resp.status_code == 403
        assert "Access denied" in rogue_dl_resp.text

        # Admin user CAN access the artifacts (admin override permitted)
        admin_headers = {"Authorization": f"Bearer {admin_token}"}
        admin_list_resp = await client.get(f"/api/v1/artifacts/{session_id}", headers=admin_headers)
        assert admin_list_resp.status_code == 200
        assert len(admin_list_resp.json()) == 2

        # Directory traversal prevention
        traversal_resp = await client.get(
            f"/api/v1/artifacts/{session_id}/../../etc/passwd",
            headers=operator_headers,
        )
        assert traversal_resp.status_code in {400, 403, 404}
