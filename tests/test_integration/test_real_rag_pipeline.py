"""
tests/test_integration/test_real_rag_pipeline.py
================================================
Integration tests verifying genuine document parsing, chunking, embedding,
persistence, and vector retrieval with source-level citations.
"""

from pathlib import Path

import docx
import openpyxl
import pytest
from app.knowledge.service import KnowledgeService
from PIL import Image, ImageDraw
from sqlalchemy.ext.asyncio import AsyncSession

from tests.test_integration.helpers import get_available_ollama_models, is_ollama_reachable

pytestmark = pytest.mark.skipif(
    not is_ollama_reachable(),
    reason="Local Ollama daemon is not running on 127.0.0.1:11434",
)


@pytest.fixture
def sample_docx_file(tmp_path: Path) -> Path:
    p = tmp_path / "boiler_maintenance_manual.docx"
    doc = docx.Document()
    doc.add_heading("Section 1: High Pressure Boiler Maintenance", level=1)
    doc.add_paragraph(
        "Operating pressure on steam drum B-201 must never exceed 18.5 bar gauge. "
        "The safety relief valve PSV-102 is calibrated to lift at 19.0 bar."
    )
    doc.add_heading("Section 2: Emergency Shutdown Protocol", level=1)
    doc.add_paragraph(
        "In the event of flameout, initiate fuel gas trip valve closure within 1.5 seconds. "
        "Purge firebox with nitrogen before any restart attempt."
    )
    doc.save(str(p))
    return p


@pytest.fixture
def sample_xlsx_file(tmp_path: Path) -> Path:
    p = tmp_path / "pump_operating_telemetry.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Telemetry"
    ws.append(["Pump_ID", "Flow_Rate_GPM", "Discharge_Bar", "Vibration_RMS"])
    ws.append(["P-101A", 450.5, 12.4, 2.1])
    ws.append(["P-101B", 448.0, 12.1, 4.8])
    ws.append(["P-102A", 310.2, 8.5, 1.8])
    wb.save(str(p))
    return p


@pytest.fixture
def sample_image_file(tmp_path: Path) -> Path:
    p = tmp_path / "sensor_tag_label.png"
    img = Image.new("RGB", (300, 150), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((10, 20), "SENSOR TAG: TT-401\nRANGE: -20C to 450C\nLOC: CRACKER UNIT 3", fill=(0, 0, 0))
    img.save(str(p))
    return p


@pytest.mark.asyncio
async def test_real_docx_ingestion_and_search(
    db_session: AsyncSession,
    sample_docx_file: Path,
) -> None:
    available = get_available_ollama_models()
    if not any("all-minilm" in m for m in available):
        pytest.skip("all-minilm:l6-v2 model not available")

    svc = KnowledgeService()
    doc = await svc.ingest_file(
        db=db_session,
        file_path=sample_docx_file,
        filename=sample_docx_file.name,
        collection_id="integration_test",
    )

    assert doc.id is not None
    assert doc.total_chunks >= 2
    assert doc.filename == "boiler_maintenance_manual.docx"

    # Search for drum pressure limits
    results = await svc.search(
        db=db_session,
        query="steam drum maximum operating pressure limit",
        collection_id="integration_test",
        top_k=3,
    )

    assert len(results) > 0
    top = results[0]
    assert "18.5 bar" in top.content or "PSV-102" in top.content
    assert top.filename == "boiler_maintenance_manual.docx"
    assert "Source" in top.citation_header()


@pytest.mark.asyncio
async def test_real_xlsx_ingestion_and_search(
    db_session: AsyncSession,
    sample_xlsx_file: Path,
) -> None:
    available = get_available_ollama_models()
    if not any("all-minilm" in m for m in available):
        pytest.skip("all-minilm:l6-v2 model not available")

    svc = KnowledgeService()
    doc = await svc.ingest_file(
        db=db_session,
        file_path=sample_xlsx_file,
        filename=sample_xlsx_file.name,
        collection_id="integration_test",
    )

    assert doc.total_chunks >= 1

    # Search for high vibration pump
    results = await svc.search(
        db=db_session,
        query="pump vibration telemetry flow rate",
        collection_id="integration_test",
        top_k=3,
    )

    assert len(results) > 0
    assert any("P-101B" in r.content or "Vibration" in r.content for r in results)


@pytest.mark.asyncio
async def test_real_image_ingestion_and_search(
    db_session: AsyncSession,
    sample_image_file: Path,
) -> None:
    available = get_available_ollama_models()
    if not any("all-minilm" in m for m in available):
        pytest.skip("all-minilm:l6-v2 model not available")

    svc = KnowledgeService()
    doc = await svc.ingest_file(
        db=db_session,
        file_path=sample_image_file,
        filename=sample_image_file.name,
        collection_id="integration_test",
    )

    assert doc.total_chunks >= 1

    results = await svc.search(
        db=db_session,
        query="sensor tag visual asset image",
        collection_id="integration_test",
        top_k=2,
    )

    assert len(results) > 0
    assert results[0].filename == "sensor_tag_label.png"
