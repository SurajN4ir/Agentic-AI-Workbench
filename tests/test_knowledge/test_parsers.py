"""
tests/test_knowledge/test_parsers.py
====================================
Tests for layout-aware document parsers (PDF, Word, CSV, XLSX, Text).
"""

from pathlib import Path

import openpyxl
import pytest
from app.knowledge.parsers.docx_parser import DocxParser
from app.knowledge.parsers.spreadsheet_parser import SpreadsheetParser
from app.knowledge.parsers.text_parser import TextParser
from app.knowledge.parsers.unified import UnifiedDocumentParser, compute_file_sha256
from docx import Document


class TestDocumentParsers:
    def test_docx_parser_extracts_headings_paragraphs_tables(self, tmp_path: Path) -> None:
        """Verifies DOCX parser extracts paragraphs and converts tables to Markdown."""
        doc_path = tmp_path / "test_spec.docx"
        doc = Document()
        doc.add_heading("System Architecture", level=1)
        doc.add_paragraph("This document outlines the sovereign air-gapped system.")

        # Add a 2x3 table
        table = doc.add_table(rows=2, cols=3)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Component"
        hdr_cells[1].text = "Technology"
        hdr_cells[2].text = "Status"

        row_cells = table.rows[1].cells
        row_cells[0].text = "Backend"
        row_cells[1].text = "FastAPI"
        row_cells[2].text = "Operational"

        doc.save(str(doc_path))

        parser = DocxParser()
        sections = parser.extract_sections(doc_path)

        assert len(sections) >= 3
        types = [s.section_type for s in sections]
        assert "heading" in types
        assert "paragraph" in types
        assert "table" in types

        table_sec = next(s for s in sections if s.section_type == "table")
        assert "| Component | Technology | Status |" in table_sec.text
        assert "| Backend | FastAPI | Operational |" in table_sec.text

    def test_spreadsheet_parser_csv(self, tmp_path: Path) -> None:
        """Verifies CSV parsing formats rows as Markdown tables."""
        csv_path = tmp_path / "telemetry.csv"
        csv_content = "timestamp,sensor_id,temperature,pressure\n2026-10-17T00:00:00,S-01,24.5,101.3\n2026-10-17T00:01:00,S-01,24.8,101.5\n"
        csv_path.write_text(csv_content, encoding="utf-8")

        parser = SpreadsheetParser()
        sections = parser.extract_sections(csv_path)

        assert len(sections) >= 1
        assert sections[0].section_type == "table"
        assert "| timestamp | sensor_id | temperature | pressure |" in sections[0].text
        assert "| 2026-10-17T00:00:00 | S-01 | 24.5 | 101.3 |" in sections[0].text

    def test_spreadsheet_parser_xlsx(self, tmp_path: Path) -> None:
        """Verifies XLSX parsing formats sheets as Markdown tables."""
        xlsx_path = tmp_path / "inventory.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Assets"
        ws.append(["Asset ID", "Location", "Criticality"])
        ws.append(["A-100", "Turbine Room", "High"])
        wb.save(xlsx_path)

        parser = SpreadsheetParser()
        sections = parser.extract_sections(xlsx_path)

        assert len(sections) >= 1
        assert sections[0].section_type == "table"
        assert "### Sheet: Assets" in sections[0].text
        assert "| Asset ID | Location | Criticality |" in sections[0].text
        assert "| A-100 | Turbine Room | High |" in sections[0].text

    def test_text_parser_markdown_and_code(self, tmp_path: Path) -> None:
        """Verifies text and markdown splitting into heading and paragraph blocks."""
        txt_path = tmp_path / "guide.md"
        content = "# Installation Guide\n\nEnsure Python 3.11 is installed.\n\nRun the test suite via pytest."
        txt_path.write_text(content, encoding="utf-8")

        parser = TextParser()
        sections = parser.extract_sections(txt_path)

        assert len(sections) == 3
        assert sections[0].section_type == "heading"
        assert sections[0].text == "# Installation Guide"
        assert sections[1].section_type == "paragraph"
        assert "Ensure Python 3.11" in sections[1].text

    def test_unified_parser_routing_and_hash(self, tmp_path: Path) -> None:
        """Unified parser dispatches properly and computes cryptographic SHA-256."""
        f_path = tmp_path / "notes.txt"
        f_path.write_text("Confidential industrial procedure notes.", encoding="utf-8")

        parser = UnifiedDocumentParser()
        file_hash, sections = parser.parse(f_path)

        assert len(file_hash) == 64  # Valid SHA-256 hex string
        assert file_hash == compute_file_sha256(f_path)
        assert len(sections) == 1
        assert "Confidential industrial procedure notes." in sections[0].text

    def test_unified_parser_missing_file_raises_not_found(self, tmp_path: Path) -> None:
        parser = UnifiedDocumentParser()
        with pytest.raises(FileNotFoundError):
            parser.parse(tmp_path / "nonexistent.pdf")
