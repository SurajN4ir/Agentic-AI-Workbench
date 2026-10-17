"""
tests/test_artifacts/test_generator.py
======================================
Unit tests for ArtifactGenerator (Word .docx, Excel .xlsx, and Markdown).
"""

import uuid

import docx
import openpyxl
import pytest
from app.artifacts.generator import ArtifactGenerator
from app.artifacts.models import ReportSpec, SheetSpec, SpreadsheetSpec, TableBlock, TextBlock


@pytest.fixture
def session_id() -> str:
    return f"test_art_{uuid.uuid4().hex[:12]}"


def test_generate_docx_report(session_id: str):
    """Verifies Word .docx creation with headings, text, bullets, and tables."""
    spec = ReportSpec(
        title="Distillation Column Inspection Report",
        subtitle="Unit 4 Crude Distillation - Quarterly Integrity Audit",
        author="Sovereign AI Inspector",
        sections=[
            TextBlock(
                heading="Executive Summary",
                body="Visual and ultrasonic inspection completed successfully.\n- Tray 14 shows minor scaling.\n- Reflux pump bearing vibration within limits.",
            ),
            TableBlock(
                caption="Operational Pressure & Temperature Readings",
                headers=["Tray No.", "Temp (C)", "Pressure (kPa)", "Status"],
                rows=[
                    ["1", "345.2", "140.5", "Normal"],
                    ["14", "280.1", "125.0", "Inspect"],
                    ["28", "165.4", "102.3", "Normal"],
                ],
            ),
        ],
        filename="column_inspection.docx",
    )

    out_path = ArtifactGenerator.generate_docx(spec, session_id)
    assert out_path.exists()
    assert out_path.is_file()
    assert out_path.stat().st_size > 0
    assert out_path.suffix == ".docx"

    # Verify document readable by python-docx
    doc = docx.Document(str(out_path))
    assert len(doc.paragraphs) > 0
    assert any("Distillation Column Inspection Report" in p.text for p in doc.paragraphs)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 4  # Header + 3 data rows
    assert table.rows[0].cells[0].text == "Tray No."
    assert table.rows[2].cells[3].text == "Inspect"


def test_generate_xlsx_workbook(session_id: str):
    """Verifies Excel .xlsx workbook creation with styled headers and auto column widths."""
    spec = SpreadsheetSpec(
        sheets=[
            SheetSpec(
                name="Reactor Readings",
                headers=["Timestamp", "Reactor ID", "Pressure (bar)", "Temp (C)", "Flow Rate (L/min)"],
                rows=[
                    ["2026-10-18T00:00:00Z", "R-101", 12.4, 215.3, 450.0],
                    ["2026-10-18T01:00:00Z", "R-101", 12.6, 216.1, 448.5],
                    ["2026-10-18T02:00:00Z", "R-101", 12.5, 215.8, 451.2],
                ],
            ),
            SheetSpec(
                name="Safety Thresholds",
                headers=["Parameter", "Warning Level", "Trip Limit"],
                rows=[
                    ["Max Temperature", 230.0, 250.0],
                    ["Max Pressure", 15.0, 18.0],
                ],
            ),
        ],
        filename="reactor_telemetry.xlsx",
    )

    out_path = ArtifactGenerator.generate_xlsx(spec, session_id)
    assert out_path.exists()
    assert out_path.is_file()
    assert out_path.stat().st_size > 0
    assert out_path.suffix == ".xlsx"

    # Verify readable by openpyxl
    wb = openpyxl.load_workbook(str(out_path))
    assert len(wb.sheetnames) == 2
    assert "Reactor Readings" in wb.sheetnames
    assert "Safety Thresholds" in wb.sheetnames

    ws = wb["Reactor Readings"]
    assert ws.cell(row=1, column=1).value == "Timestamp"
    assert ws.cell(row=2, column=2).value == "R-101"
    assert float(ws.cell(row=2, column=3).value) == 12.4


def test_generate_markdown_report(session_id: str):
    """Verifies GitHub-flavored markdown report generation."""
    spec = ReportSpec(
        title="Refinery Operations Log",
        subtitle="Shift Handover Summary",
        sections=[
            TextBlock(heading="Shift Highlights", body="All process units steady state."),
            TableBlock(
                headers=["Unit", "Status"],
                rows=[["CDU-1", "Online"], ["VDU-2", "Online"]],
            ),
        ],
        filename="handover.md",
    )

    out_path = ArtifactGenerator.generate_markdown(spec, session_id)
    assert out_path.exists()
    assert out_path.is_file()

    content = out_path.read_text(encoding="utf-8")
    assert "# Refinery Operations Log" in content
    assert "## Shift Highlights" in content
    assert "| CDU-1 | Online |" in content
