"""Regression tests for the agent-facing deliverable schema and handler contract."""

from pathlib import Path

import docx
import pytest
from app.agent.graph import AgentGraphBuilder
from app.tools.doc_gen_validation import normalize_doc_gen_input
from app.tools.interface import ToolStatus, validate_workspace_path
from app.tools.registry import ToolRegistry


def test_graph_accepts_section_based_docx_and_rejects_empty_body() -> None:
    graph = AgentGraphBuilder.__new__(AgentGraphBuilder)
    valid = {
        "format": "docx",
        "title": "Inspection Summary",
        "sections": [{"heading": "Finding", "body": "Casing torque is 450 Nm."}],
    }
    assert graph._validate_tool_arguments("doc_gen", valid) is None
    assert "sections[0].body" in str(
        graph._validate_tool_arguments(
            "doc_gen",
            {"format": "docx", "title": "Inspection Summary", "sections": [{"body": "  "}]},
        )
    )


def test_doc_gen_rejects_empty_or_malformed_arguments() -> None:
    cases = [
        (None, "JSON object"),
        ({"format": "docx", "title": "Title"}, "requires non-empty section content"),
        ({"format": "docx", "title": "Title", "content": "  "}, "requires non-empty section content"),
        ({"format": "docx", "title": "Title", "sections": "wrong"}, "sections' must be a list"),
        ({"format": "xlsx", "title": "Sheet", "table": {"headers": ["A"], "rows": []}}, "at least one data row"),
        (
            {"format": "xlsx", "title": "Sheet", "table": {"headers": ["A", "B"], "rows": [[1]]}},
            "exactly 2 cells",
        ),
        ({"format": "pdf", "title": "Title", "content": "Text"}, "unsupported doc_gen format"),
    ]
    for arguments, expected_error in cases:
        normalized, error = normalize_doc_gen_input(arguments)
        assert normalized is None
        assert error is not None and expected_error in error


def test_doc_gen_generates_parseable_docx_from_sections(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import app.artifacts.generator as generator

    real_validate = validate_workspace_path
    monkeypatch.setattr(
        generator,
        "validate_workspace_path",
        lambda path, session_id, create_parents=True: real_validate(
            path, session_id, create_parents=create_parents, base_override=tmp_path
        ),
    )
    registry = ToolRegistry()
    result = registry._execute_doc_gen(
        {
            "format": "docx",
            "title": "Synthetic Torque Verification",
            "filename": "browser_demo_report.docx",
            "sections": [
                {
                    "heading": "Verified finding",
                    "body": "The synthetic operating manual specifies casing torque of 450 Nm.",
                }
            ],
        },
        "docgen_test_session",
    )

    assert result.status == ToolStatus.SUCCESS
    generated_path = Path(result.data["generated_files"][0])
    assert generated_path.is_file() and generated_path.stat().st_size > 0
    parsed = docx.Document(generated_path)
    parsed_text = "\n".join(paragraph.text for paragraph in parsed.paragraphs)
    assert "Synthetic Torque Verification" in parsed_text
    assert "450 Nm" in parsed_text


def test_doc_gen_returns_tool_error_without_creating_title_only_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import app.artifacts.generator as generator

    real_validate = validate_workspace_path
    monkeypatch.setattr(
        generator,
        "validate_workspace_path",
        lambda path, session_id, create_parents=True: real_validate(
            path, session_id, create_parents=create_parents, base_override=tmp_path
        ),
    )
    result = ToolRegistry()._execute_doc_gen(
        {"format": "docx", "title": "Must Not Be Empty", "sections": []},
        "empty_doc_session",
    )

    assert result.status == ToolStatus.ERROR
    assert "requires non-empty section content" in str(result.error)
    assert not (tmp_path / "empty_doc_session").exists()


def test_doc_gen_supports_legacy_content_alias(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import app.artifacts.generator as generator

    real_validate = validate_workspace_path
    monkeypatch.setattr(
        generator,
        "validate_workspace_path",
        lambda path, session_id, create_parents=True: real_validate(
            path, session_id, create_parents=create_parents, base_override=tmp_path
        ),
    )
    result = ToolRegistry()._execute_doc_gen(
        {"format": "docx", "title": "Legacy Request", "content": "Legacy body remains supported."},
        "legacy_doc_session",
    )
    assert result.status == ToolStatus.SUCCESS
    assert "Legacy body remains supported." in "\n".join(
        paragraph.text for paragraph in docx.Document(result.data["generated_files"][0]).paragraphs
    )
