"""
app/knowledge/parsers/spreadsheet_parser.py
===========================================
Parses CSV and Excel (.xlsx) spreadsheets into structured Markdown tables.
"""

import csv
import logging
from pathlib import Path

import openpyxl

from app.knowledge.parsers.base import BaseDocumentParser, DocumentSection

logger = logging.getLogger("sovereign.knowledge.spreadsheet_parser")


class SpreadsheetParser(BaseDocumentParser):
    """Extracts tables and records from CSV and XLSX files."""

    def extract_sections(self, file_path: Path) -> list[DocumentSection]:
        ext = file_path.suffix.lower()
        if ext == ".csv":
            return self._parse_csv(file_path)
        elif ext == ".xlsx":
            return self._parse_xlsx(file_path)
        return []

    def _parse_csv(self, file_path: Path) -> list[DocumentSection]:
        sections: list[DocumentSection] = []
        try:
            with open(file_path, encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                rows = [row for row in reader if any(cell.strip() for cell in row)]
        except Exception as e:
            logger.error("Failed to read CSV %s: %s", file_path, e)
            return []

        if not rows:
            return []

        col_count = len(rows[0])
        header = "| " + " | ".join(rows[0]) + " |"
        divider = "| " + " | ".join(["---"] * col_count) + " |"

        # Split large CSVs into chunks of 30 rows with repeated header
        batch_size = 30
        for i in range(1, len(rows), batch_size):
            chunk_rows = rows[i : i + batch_size]
            body = ["| " + " | ".join(row) + " |" for row in chunk_rows]
            table_md = "\n".join([header, divider] + body)
            sections.append(
                DocumentSection(
                    text=table_md,
                    section_type="table",
                    metadata={"row_start": i, "row_end": min(i + batch_size, len(rows))},
                )
            )

        return sections

    def _parse_xlsx(self, file_path: Path) -> list[DocumentSection]:
        sections: list[DocumentSection] = []
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        except Exception as e:
            logger.error("Failed to open XLSX %s: %s", file_path, e)
            return []

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            raw_rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                str_cells = [str(cell).strip() if cell is not None else "" for cell in row]
                if any(str_cells):
                    raw_rows.append(str_cells)

            if not raw_rows:
                continue

            col_count = len(raw_rows[0])
            header = f"### Sheet: {sheet_name}\n| " + " | ".join(raw_rows[0]) + " |"
            divider = "| " + " | ".join(["---"] * col_count) + " |"

            batch_size = 30
            for i in range(1, len(raw_rows), batch_size):
                chunk_rows = raw_rows[i : i + batch_size]
                body = ["| " + " | ".join(row) + " |" for row in chunk_rows]
                table_md = "\n".join([header, divider] + body)
                sections.append(
                    DocumentSection(
                        text=table_md,
                        section_type="table",
                        metadata={"sheet": sheet_name, "row_start": i, "row_end": min(i + batch_size, len(raw_rows))},
                    )
                )

        return sections
